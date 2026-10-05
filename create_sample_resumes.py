import os
from docx import Document

os.makedirs("sample_resumes", exist_ok=True)

# 1. Data Scientist Resume
doc1 = Document()
doc1.add_heading("Alex Rivera - Machine Learning Engineer", 0)
doc1.add_paragraph("Email: alex.rivera@example.com | Phone: +1 555-0199 | Location: San Francisco, CA")
doc1.add_paragraph("LinkedIn: linkedin.com/in/alexrivera-ml | GitHub: github.com/alexrivera")

doc1.add_heading("Professional Summary", level=1)
doc1.add_paragraph(
    "Data Scientist and Machine Learning Engineer with 3+ years of experience designing and deploying "
    "end-to-end ML and NLP pipelines. Proficient in Python, PyTorch, TensorFlow, Scikit-learn, SQL, "
    "and cloud deployments on AWS. Strong background in statistical modeling, feature engineering, and LLM fine-tuning."
)

doc1.add_heading("Technical Skills", level=1)
doc1.add_paragraph(
    "Programming: Python, SQL, R, Bash\n"
    "ML & AI: PyTorch, TensorFlow, Scikit-learn, Hugging Face, Transformers, NLP, Computer Vision\n"
    "Data & Tools: Pandas, NumPy, Spark, Docker, Git, MLflow, Airflow\n"
    "Cloud & DB: AWS (S3, SageMaker, EC2), PostgreSQL, MongoDB"
)

doc1.add_heading("Work Experience", level=1)
doc1.add_heading("Senior AI Engineer | DataTech Solutions (2022 - Present)", level=2)
doc1.add_paragraph(
    "- Designed and deployed a transformer-based semantic search and recommendation engine serving 500k+ daily queries.\n"
    "- Built automated ETL and feature engineering pipelines using Python, Pandas, and Apache Spark.\n"
    "- Improved recommendation model F1-score by 18% through hyperparameter tuning and custom embeddings."
)

doc1.add_heading("Machine Learning Intern | AI Innovators (2021 - 2022)", level=2)
doc1.add_paragraph(
    "- Implemented natural language processing algorithms for sentiment classification and document summarization.\n"
    "- Trained Scikit-learn classification models (Random Forest, Logistic Regression, XGBoost) achieving 92% accuracy."
)

doc1.add_heading("Education", level=1)
doc1.add_paragraph("Master of Science in Computer Science (Specialization in AI/ML)\nStanford University (2019 - 2021)")
doc1.add_paragraph("Bachelor of Technology in Information Technology\nABC University (2015 - 2019)")

doc1.save("sample_resumes/sample_data_scientist.docx")

# 2. Software / Full Stack Engineer Resume
doc2 = Document()
doc2.add_heading("Priya Sharma - Full Stack Developer", 0)
doc2.add_paragraph("Email: priya.sharma@example.com | Phone: +1 555-0144 | Location: Austin, TX")
doc2.add_paragraph("GitHub: github.com/priyasharma-dev | Portfolio: priyasharma.dev")

doc2.add_heading("Professional Summary", level=1)
doc2.add_paragraph(
    "Full Stack Software Engineer with 4 years of experience building scalable web applications, REST APIs, "
    "and microservices using React, Node.js, TypeScript, Python, and Docker. Passionate about clean code, CI/CD, "
    "and high-performance backend architectures."
)

doc2.add_heading("Technical Skills", level=1)
doc2.add_paragraph(
    "Frontend: React, Next.js, TypeScript, JavaScript, HTML5, CSS3, TailwindCSS\n"
    "Backend: Node.js, Express, Python, FastAPI, Django, RESTful APIs, GraphQL\n"
    "Databases & Cloud: PostgreSQL, MongoDB, Redis, Docker, Kubernetes, AWS, CI/CD, Git"
)

doc2.add_heading("Work Experience", level=1)
doc2.add_heading("Software Engineer | CloudScale Systems (2022 - Present)", level=2)
doc2.add_paragraph(
    "- Developed full-stack SaaS features using React, TypeScript, Node.js, and PostgreSQL.\n"
    "- Optimized backend API endpoints reducing latency by 40% with Redis caching.\n"
    "- Containerized microservices using Docker and orchestrated deployments on Kubernetes."
)

doc2.add_heading("Education", level=1)
doc2.add_paragraph("Bachelor of Science in Computer Science\nUniversity of Texas at Austin (2018 - 2022)")

doc2.save("sample_resumes/sample_software_engineer.docx")

print("Generated sample resumes in sample_resumes/")
